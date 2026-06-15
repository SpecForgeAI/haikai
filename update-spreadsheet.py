"""Update the Excel spreadsheet with newly-created service IDs + canonical pack IDs + subfolder + status."""
import openpyxl, shutil
from openpyxl.styles import PatternFill, Font

XLSX = r'C:\Users\gazwa\OneDrive\Documents\SDD_App\discovery_packs.xlsx'

# Inline results from the successful run (script output captured above)
CREATED = {
    'PetClinic':              {'svc_id': 'svc-mob4f5ai-h4yfj', 'app_id': 'app-moui8j1z-lo1sn', 'comp_id': 'comp-moz5c2nc-6wye2', 'lang_pack': 'java-lang',       'fw_pack': 'java-spring-boot',  'subfolder': '',                       'filename': 'test-repo-petclinic.json'},
    'nestjs-realworld':       {'svc_id': 'svc-mo8nsfj7-57h4r', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'typescript-lang', 'fw_pack': 'nestjs',            'subfolder': '',                       'filename': 'test-repo-nestjs-realworld.json'},
    'Saleor':                 {'svc_id': 'svc-mosqw9pa-5jlxc', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'python-lang',     'fw_pack': 'django',            'subfolder': 'saleor',                 'filename': 'test-repo-saleor.json'},
    'react-redux-realworld':  {'svc_id': 'svc-modrxy6b-4prqe', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'javascript-lang', 'fw_pack': 'react-javascript',  'subfolder': '',                       'filename': 'test-repo-react-redux-realworld.json'},
    'angular-realworld':      {'svc_id': 'svc-moyv6pl2-hxrbg', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'typescript-lang', 'fw_pack': 'angular',           'subfolder': '',                       'filename': 'test-repo-angular-realworld.json'},
    'WordPress':              {'svc_id': 'svc-moat53u0-gfnwb', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'php-lang',        'fw_pack': 'wordpress',         'subfolder': '',                       'filename': 'test-repo-wordpress.json'},
    'eShopOnWeb':             {'svc_id': 'svc-mo7dpzyf-o4lac', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'csharp-lang',     'fw_pack': 'asp-net-core',      'subfolder': 'src/Web',                'filename': 'test-repo-eshoponweb.json'},
    'flask-microblog':        {'svc_id': 'svc-mo2c3k9k-bdeas', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'python-lang',     'fw_pack': 'flask',             'subfolder': '',                       'filename': 'test-repo-flask-microblog.json'},
    'Redmine':                {'svc_id': 'svc-mowyibpn-fecnk', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'ruby-lang',       'fw_pack': 'rails',             'subfolder': '',                       'filename': 'test-repo-redmine.json'},
    'Discourse':              {'svc_id': 'svc-mohrmmqz-ly0yn', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'ruby-lang',       'fw_pack': 'rails',             'subfolder': '',                       'filename': 'test-repo-discourse.json'},
    'beer-shop-go':           {'svc_id': 'svc-moas6iuv-hh084', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'go-lang',         'fw_pack': 'kratos',            'subfolder': 'app',                    'filename': 'test-repo-beer-shop-go.json'},
    'OrangeHRM':              {'svc_id': 'svc-mo2fwkti-2tk93', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'php-lang',        'fw_pack': 'symfony',           'subfolder': 'src',                    'filename': 'test-repo-orangehrm.json'},
    'eShopLegacyMVC':         {'svc_id': 'svc-mofvde8z-j25k3', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'csharp-lang',     'fw_pack': 'asp-net-framework', 'subfolder': 'eShopLegacyMVCSolution', 'filename': 'test-repo-eshoplegacymvc.json'},
    'magento-lts':            {'svc_id': 'svc-moircdjc-oy1nl', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'php-lang',        'fw_pack': 'magento',           'subfolder': '',                       'filename': 'test-repo-magento-lts.json'},
    'jquery-ui':              {'svc_id': 'svc-moaequpm-igk0k', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'javascript-lang', 'fw_pack': 'jquery',            'subfolder': 'ui',                     'filename': 'test-repo-jquery-ui.json'},
    'wxWidgets':              {'svc_id': 'svc-movs160z-hb5wg', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'cpp-lang',        'fw_pack': 'wxwidgets',         'subfolder': 'src',                    'filename': 'test-repo-wxwidgets.json'},
    'oatpp-crud':             {'svc_id': 'svc-mo0c8iab-58e5a', 'app_id': '?',                   'comp_id': '?',                  'lang_pack': 'cpp-lang',        'fw_pack': 'oatpp',             'subfolder': 'src',                    'filename': 'test-repo-oatpp-crud.json'},
}

SKIPPED = {
    'sunflower':                'No kotlin language pack registered',
    'cobol-programming-course': 'No cobol language pack registered',
    'redis':                    'No C language pack registered',
    'systemd':                  'No C language pack registered',
}

ALREADY_DONE = {'Scenarios Frontend', 'Scenarios Service', 'OpenMRS'}


def main():
    try:
        wb = openpyxl.load_workbook(XLSX)
    except PermissionError:
        print(f'ERROR: cannot open {XLSX} -- close Excel first.')
        return 1
    ws = wb['Test Repos Projects']

    if ws.cell(row=1, column=8).value != 'Repo Subfolder':
        ws.cell(row=1, column=8, value='Repo Subfolder').font = Font(bold=True)
        ws.cell(row=1, column=9, value='Status').font = Font(bold=True)
        ws.cell(row=1, column=10, value='Application Id').font = Font(bold=True)
        ws.cell(row=1, column=11, value='App Component Id').font = Font(bold=True)
        ws.cell(row=1, column=12, value='Filename').font = Font(bold=True)

    GREEN = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
    GREY = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')
    YELLOW = PatternFill(start_color='FFF3CD', end_color='FFF3CD', fill_type='solid')

    updated = skipped = already = 0
    for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
        if r_idx == 1:
            continue
        name = row[0].value
        if not name:
            continue
        if name in ALREADY_DONE:
            already += 1
            continue
        if name in SKIPPED:
            ws.cell(row=r_idx, column=9, value='Skipped (no pack)').fill = GREY
            existing = ws.cell(row=r_idx, column=7).value or ''
            ws.cell(row=r_idx, column=7, value=(existing + ' | ' + SKIPPED[name]).strip(' |'))
            skipped += 1
            continue
        if name in CREATED:
            c = CREATED[name]
            ws.cell(row=r_idx, column=3, value=c['lang_pack'])
            ws.cell(row=r_idx, column=4, value=c['fw_pack'])
            ws.cell(row=r_idx, column=6, value=c['svc_id'])
            ws.cell(row=r_idx, column=8, value=c.get('subfolder', '') or '')
            status = 'Created'
            fill = GREEN
            if name == 'PetClinic':
                status = 'Created (with leftovers)'
                fill = YELLOW
                existing = ws.cell(row=r_idx, column=7).value or ''
                note = 'NOTE: project also contains leftover entities from a prior model file (svc-mo3c2kln-fq7e8). Manual cleanup recommended via UI.'
                ws.cell(row=r_idx, column=7, value=(existing + ' | ' + note).strip(' |'))
            if name == 'eShopLegacyMVC':
                existing = ws.cell(row=r_idx, column=7).value or ''
                note = 'NOTE: original dotnet-architecture/eShopLegacyMVC was retired into dotnet-architecture/eShopModernizing. Updated repo URL accordingly; subfolder eShopLegacyMVCSolution.'
                ws.cell(row=r_idx, column=7, value=(existing + ' | ' + note).strip(' |'))
                ws.cell(row=r_idx, column=2, value='https://github.com/dotnet-architecture/eShopModernizing.git')
            ws.cell(row=r_idx, column=9, value=status).fill = fill
            ws.cell(row=r_idx, column=10, value=c['app_id'])
            ws.cell(row=r_idx, column=11, value=c['comp_id'])
            ws.cell(row=r_idx, column=12, value=c.get('filename', ''))
            updated += 1

    # Column widths
    for col_idx, width in [(6, 28), (8, 25), (9, 24), (10, 22), (11, 22), (12, 32)]:
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    backup = XLSX + '.bak'
    shutil.copyfile(XLSX, backup)
    wb.save(XLSX)
    print(f'Updated {updated} rows | Skipped {skipped} | Already-done {already}.')
    print(f'Backup: {backup}')
    return 0

if __name__ == '__main__':
    raise SystemExit(main())
