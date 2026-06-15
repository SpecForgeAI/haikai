"""One-shot helper: convert 3 worksheets in discovery_packs.xlsx into MD."""
import openpyxl, os

XLSX = r'C:\Users\gazwa\OneDrive\Documents\SDD_App\discovery_packs.xlsx'
OUT_DIR = r'C:\Workspaces\SSD\architecture-store-and-diagrams\discovery-service\perf'

XLSX_LITERAL = r'C:\Users\gazwa\OneDrive\Documents\SDD_App\discovery_packs.xlsx'


def cell(v):
    if v is None:
        return ''
    return str(v).replace('|', '\\|').replace('\n', ' ').strip()


def to_md_table(rows):
    if not rows:
        return ''
    headers = [cell(c) for c in rows[0]]
    out = ['| ' + ' | '.join(headers) + ' |',
           '|' + '|'.join(['---'] * len(headers)) + '|']
    for r in rows[1:]:
        if all(c is None or str(c).strip() == '' for c in r):
            continue
        cells = [cell(c) for c in r]
        while len(cells) < len(headers):
            cells.append('')
        cells = cells[:len(headers)]
        out.append('| ' + ' | '.join(cells) + ' |')
    return '\n'.join(out)


def load_sheet(wb, name):
    ws = wb[name]
    rows = []
    for r in range(1, ws.max_row + 1):
        rows.append([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])
    return rows


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    trp = to_md_table(load_sheet(wb, 'Test Repos Projects'))
    lp = to_md_table(load_sheet(wb, 'Language Packs'))
    fp = to_md_table(load_sheet(wb, 'Framework Packs'))

    md1 = (
        '# Test Repos — Project / Service Inventory\n\n'
        f'Source of truth for the test-repo projects + services used by the\n'
        f'discovery performance scoring system. Maintained alongside the live\n'
        f'spreadsheet at `{XLSX_LITERAL}`\n'
        '(sheet "Test Repos Projects"); regenerate this file when that sheet\n'
        'changes, or update both in lockstep via `generate-pack-mds.py`.\n\n'
        'Fields:\n\n'
        '- **Name** — short repo / service name.\n'
        '- **Repo URL** — GitHub HTTPS URL.\n'
        '- **Language Pack Name** / **Framework Pack Name** — the **canonical\n'
        '  registered pack IDs** that the discovery service uses (e.g.\n'
        '  `java-lang`, `java-spring-boot`). When `Status` is `Skipped (no pack)`\n'
        '  these columns hold the *desired* pack name that has not yet been\n'
        '  registered.\n'
        '- **Project Id** / **Service Id** — UUIDs in the architecture-model-service.\n'
        '  Used in the run-start payload (`POST /discovery/runs`).\n'
        '- **Repo Subfolder** — relative path inside the repo to scope the scan\n'
        '  (e.g. `src/Web` for monorepos). Empty = whole repo.\n'
        '- **Status** — `Created` / `Created (with leftovers)` / `Skipped (no pack)`\n'
        '  for the 17-service test-repo seeding done on 2026-04-25.\n'
        '- **Application Id** / **App Component Id** / **Filename** — internals of\n'
        '  the architecture-model-service entity, populated when the seed script\n'
        '  ran.\n\n'
        + trp + '\n'
    )

    md2 = (
        '# Language Pack Inventory (Planning)\n\n'
        f'Source: `{XLSX_LITERAL}` sheet "Language Packs". This is the **planning\n'
        'inventory** — the language packs the project intends to support, sized\n'
        'by estimated share of real-world codebase coverage.\n\n'
        '> **Important:** the `Language Pack Name` column here uses *aspirational*\n'
        '> names (e.g. `java-modern`, `python-3`, `javascript-modern`). The\n'
        '> **canonical registered pack IDs** in the discovery service codebase\n'
        '> are simpler — `java-lang`, `python-lang`, `javascript-lang`, etc.\n'
        '> When wiring services for runs, use the canonical IDs from\n'
        '> [`test-repos-projects.md`](./test-repos-projects.md) and the\n'
        '> registered packs in\n'
        '> `discovery-service/src/services/extensionPacks/languagePacks/`.\n'
        '> The `Estimated %` figures in this sheet are still useful as rough\n'
        '> coverage targets, even though several entries are not yet implemented.\n\n'
        + lp + '\n'
    )

    md3 = (
        '# Framework Pack Inventory (Planning)\n\n'
        f'Source: `{XLSX_LITERAL}` sheet "Framework Packs". The **planning\n'
        'inventory** of framework packs the project intends to support, sized\n'
        'by estimated share of real-world codebase coverage.\n\n'
        '> **Important:** the `Framework Pack Name` column uses *aspirational*\n'
        '> names (e.g. `spring-boot`, `react`). The **canonical registered pack\n'
        '> IDs** in the discovery service codebase are sometimes more specific —\n'
        '> e.g. `java-spring-boot`, `react-javascript` and `react-typescript`\n'
        '> (split per language), `spring-classic` (Hibernate-era), etc. When\n'
        '> wiring services for runs, use the canonical IDs from\n'
        '> [`test-repos-projects.md`](./test-repos-projects.md) and the\n'
        '> registered packs in\n'
        '> `discovery-service/src/services/extensionPacks/frameworkPacks/`.\n'
        '> Entries marked here that have no matching registered pack are\n'
        '> *not yet implemented*.\n\n'
        + fp + '\n'
    )

    open(os.path.join(OUT_DIR, 'test-repos-projects.md'), 'w', encoding='utf-8').write(md1)
    open(os.path.join(OUT_DIR, 'language-packs.md'), 'w', encoding='utf-8').write(md2)
    open(os.path.join(OUT_DIR, 'framework-packs.md'), 'w', encoding='utf-8').write(md3)
    print('Wrote 3 MD files in', OUT_DIR)


if __name__ == '__main__':
    main()
